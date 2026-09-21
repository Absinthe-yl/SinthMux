from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "毕业设计材料"
OUT.mkdir(exist_ok=True)

TITLE = "SinthMux跨设备共享持久终端系统设计与实现"
STUDENT = "陈珏龙"
STUDENT_ID = "3231807030"
COLLEGE = "计算机与数据科学学院"
MAJOR = "软件工程"
CLASS_NAME = "软件工程2301"
ADVISER = "林芳"
SCHOOL = "福建理工大学"

BODY_FONT = "Arial Unicode MS"
HEAD_FONT = "Arial Unicode MS"
LATIN_FONT = "Arial Unicode MS"


REFERENCES = [
    "[1] 张海藩, 牟永敏. 软件工程导论（第6版）[M]. 北京: 清华大学出版社, 2013.",
    "[2] 郑人杰, 马素霞, 殷人昆. 软件工程概论（第3版）[M]. 北京: 机械工业出版社, 2014.",
    "[3] Pressman R S, Maxim B R. Software Engineering: A Practitioner's Approach (9th Edition)[M]. New York: McGraw-Hill Education, 2019.",
    "[4] Sommerville I. Software Engineering (10th Edition)[M]. Boston: Pearson, 2015.",
    "[5] Donovan A A A, Kernighan B W. The Go Programming Language[M]. Boston: Addison-Wesley Professional, 2015.",
    "[6] Fette I, Melnikov A. RFC 6455: The WebSocket Protocol[S]. Internet Engineering Task Force, 2011.",
    "[7] Rescorla E. RFC 8446: The Transport Layer Security Protocol Version 1.3[S]. Internet Engineering Task Force, 2018.",
    "[8] Campbell B, Bradley J, Sakimura N, et al. RFC 8705: OAuth 2.0 Mutual-TLS Client Authentication and Certificate-Bound Access Tokens[S]. Internet Engineering Task Force, 2020.",
    "[9] OWASP Foundation. WebSocket Security Cheat Sheet[EB/OL]. https://cheatsheetseries.owasp.org/cheatsheets/WebSocket_Security_Cheat_Sheet.html, 2026-09-21.",
    "[10] React Team. React Documentation[EB/OL]. https://react.dev/, 2026-09-21.",
    "[11] xterm.js Contributors. xterm.js: A Terminal for the Web[EB/OL]. https://github.com/xtermjs/xterm.js, 2026-09-21.",
    "[12] tmux Contributors. tmux: Terminal Multiplexer[EB/OL]. https://github.com/tmux/tmux, 2026-09-21.",
    "[13] ShellHub Contributors. ShellHub: Centralized SSH for the Edge and Cloud Computing[EB/OL]. https://github.com/shellhub-io/shellhub, 2026-09-21.",
    "[14] Saint-Hilaire Y, Contributors. MeshCentral: Web Based Remote Computer Management Server[EB/OL]. https://github.com/Ylianst/MeshCentral, 2026-09-21.",
    "[15] Zhang E. sshx: A Secure Web-based Collaborative Terminal[EB/OL]. https://github.com/ekzhang/sshx, 2026-09-21.",
    "[16] Tao S. ttyd: Share Your Terminal over the Web[EB/OL]. https://github.com/tsl0922/ttyd, 2026-09-21.",
    "[17] PostgreSQL Global Development Group. PostgreSQL Documentation[EB/OL]. https://www.postgresql.org/docs/, 2026-09-21.",
    "[18] Caddy Authors. Caddy Documentation[EB/OL]. https://caddyserver.com/docs/, 2026-09-21.",
    "[19] Tailscale Inc. How NAT Traversal Works[EB/OL]. https://tailscale.com/blog/how-nat-traversal-works/, 2026-09-21.",
    "[20] Teleport Authors. Teleport Architecture[EB/OL]. https://goteleport.com/docs/reference/architecture/, 2026-09-21.",
]


def set_run_font(run, size=10.5, bold=False, font=BODY_FONT):
    run.font.name = LATIN_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, *, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, font=BODY_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, font=font)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    return p


def add_cell_paragraph(cell, text="", *, bold=False, size=10.5, indent=True, before=0, after=3, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.35
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_body(doc, text, *, bold=False, size=10.5, indent=True, before=0, after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True, font=HEAD_FONT)
    return p


def add_numbered(doc, label, text, *, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    r1 = p.add_run(label)
    set_run_font(r1, size=size, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=size)
    return p


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def configure_doc(doc, *, page_numbers=True):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)
    if page_numbers:
        add_page_number(section)
    normal = doc.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(0)
    return doc


def add_center_title(doc, lines):
    for text, size, bold in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, font=HEAD_FONT)


def add_label_value_table(doc, rows, widths=(3.2, 13.4)):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(widths[0])
        cells[1].width = Cm(widths[1])
        set_cell_text(cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
        set_cell_text(cells[1], value)
        prevent_row_split(table.rows[-1])
    return table


def add_signature_box(doc, title, body, date_text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.text = ""
    add_cell_paragraph(cell, title, bold=True, indent=False, size=11, after=7)
    add_cell_paragraph(cell, body, indent=True, size=10.5, after=12)
    add_cell_paragraph(cell, "签字：____________________", indent=False, size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, after=7)
    add_cell_paragraph(cell, date_text, indent=False, size=10.5, align=WD_ALIGN_PARAGRAPH.RIGHT, after=3)
    set_cell_margins(cell, top=160, start=180, bottom=160, end=180)
    return table


def build_application():
    doc = configure_doc(Document())
    add_center_title(doc, [("毕业设计课题申请表", 20, True)])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(f"学生：{STUDENT}    学号：{STUDENT_ID}    班级：{CLASS_NAME}")
    set_run_font(r, 10.5)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Cm(3.2)
    table.columns[1].width = Cm(13.4)

    rows = [
        ("课题名称", TITLE),
        ("课题简介", None),
        ("课题要求\n业务功能\n技术指标", None),
        ("参考资料", None),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
        if value is not None:
            set_cell_text(cells[1], value, bold=True, size=11)
            prevent_row_split(table.rows[-1])
        else:
            cells[1].text = ""
            set_cell_margins(cells[1], top=130, start=160, bottom=130, end=160)

    cell = table.rows[1].cells[1]
    add_cell_paragraph(cell, "SinthMux面向个人与小团队的多主机远程终端管理需求，拟实现一个由公网Hub、远程Agent和浏览器端组成的跨设备共享持久终端平台。用户在个人电脑、实验室主机或云服务器上安装Agent，Agent通过加密的反向长连接主动接入Hub，无需为受控主机配置公网IP、端口映射或暴露SSH端口。用户可从电脑、手机和平板浏览器查看主机状态，创建、恢复、切换和关闭远程tmux会话，并在浏览器关闭或网络短时中断后继续原有任务。", indent=True)
    add_cell_paragraph(cell, "课题重点研究多主机设备注册与路由、终端双向流中继、持久会话管理、断线检测与自动恢复、短期票据、设备身份认证和操作审计等问题。系统采用Hub集中控制、Agent主动出站的架构，在完成可运行产品的同时，通过接入成功率、终端交互时延、断线恢复时间、并发资源占用和越权拦截率等指标评价方案的有效性。", indent=True)

    cell = table.rows[2].cells[1]
    add_cell_paragraph(cell, "一、业务功能", bold=True, indent=False, size=11)
    functional = [
        "实现用户登录、设备配对、设备列表、在线状态和最后在线时间展示。",
        "实现Agent主动连接Hub、心跳上报、自动重连和能力上报，支持Linux与macOS，Windows首版通过WSL接入。",
        "实现远程tmux会话的查询、创建、重命名、附着、窗口尺寸同步和关闭。",
        "实现基于xterm.js的Web终端，支持电脑、手机和平板浏览器访问，并提供移动端常用特殊键。",
        "实现浏览器断线后的重新签票、自动重连和同一tmux会话恢复，保证后台任务持续运行。",
        "实现Owner、Admin、Operator、Viewer角色及会话查看、输入、创建、关闭等细粒度权限。",
        "实现关键操作审计、设备吊销和会话连接记录；增强阶段实现多窗格工作区、受限文件管理和限时分享。",
    ]
    for i, item in enumerate(functional, 1):
        add_cell_paragraph(cell, f"{i}）{item}", indent=False, after=2)
    add_cell_paragraph(cell, "二、技术指标", bold=True, indent=False, size=11, before=4)
    metrics = [
        "采用Go实现Hub和Agent，React、TypeScript、TanStack Query与xterm.js实现Web/PWA，PostgreSQL实现持久化。",
        "Agent仅建立到Hub TCP 443端口的出站TLS连接，不要求受控主机开放新增入站端口。",
        "设备配对采用5分钟有效的一次性配对码；正式设备连接采用mTLS，浏览器终端连接采用30至60秒有效且只能兑换一次的票据。",
        "Agent每15秒发送心跳，Hub连续45秒未收到心跳则将设备标记离线；网络恢复后Agent采用指数退避自动重连。",
        "系统应支持不少于3台异构主机同时在线、每台不少于5个tmux会话，并完成不少于20路并发终端连接测试。",
        "在稳定网络环境下，终端按键到回显的P95附加时延目标不高于200毫秒；短时断网恢复时间目标不高于10秒。",
        "所有REST请求与WebSocket操作均执行服务端授权检查，关键安全测试中越权请求拦截率达到100%。",
        "按照软件工程规范完成需求、设计、编码、测试、部署文档和论文，提交可运行软件、源代码、测试报告及演示材料。",
    ]
    for i, item in enumerate(metrics, 1):
        add_cell_paragraph(cell, f"{i}）{item}", indent=False, after=2)

    cell = table.rows[3].cells[1]
    for ref in REFERENCES:
        add_cell_paragraph(cell, ref, indent=False, size=9.3, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)

    path = OUT / f"0-{STUDENT_ID}-{STUDENT}-选题申请表-SinthMux.docx"
    doc.save(path)
    return path


def build_task_book():
    doc = configure_doc(Document())
    add_center_title(doc, [(SCHOOL, 18, True), ("本科毕业设计论文任务书", 20, True)])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("填表时间：2026年12月20日")
    set_run_font(r, 10.5)
    add_body(doc, "说明：毕业设计任务书由指导教师根据课题具体情况填写，经教研室主任审批签字后生效，并作为学生开展毕业设计和成果验收的重要依据。", indent=False, size=9.5, after=10)

    info = add_label_value_table(doc, [
        ("学生姓名", STUDENT),
        ("学号", STUDENT_ID),
        ("学院", COLLEGE),
        ("专业班级", CLASS_NAME),
        ("指导教师", ADVISER),
        ("设计题目", TITLE),
    ])
    info.rows[-1].cells[1].paragraphs[0].runs[0].font.bold = True

    add_heading(doc, "一 功能要求", 1)
    add_body(doc, "本课题拟设计并实现SinthMux跨设备共享持久终端系统。系统由公网Hub、安装在远程主机上的Agent和浏览器Web/PWA组成。Agent主动建立到Hub的加密长连接，Hub维护设备在线注册表并中继控制消息和终端数据，用户无需为个人电脑或云服务器配置公网IP和端口映射，即可通过电脑、手机或平板接续远程tmux会话。", indent=True)
    requirements = [
        "用户与设备管理：实现账号登录、用户空间、一次性配对码、设备注册、在线状态、最后在线时间和设备吊销。",
        "Agent连接管理：实现Agent主动出站连接、版本与能力上报、15秒心跳、离线检测和指数退避重连。",
        "持久会话管理：实现远程tmux会话列表、创建、重命名、附着、窗口尺寸同步、关闭和断线后恢复。",
        "Web终端：采用xterm.js实现终端显示和输入，兼容桌面、手机和平板，支持移动端常用控制键和文本发送。",
        "安全控制：实现mTLS设备身份、一次性终端票据、角色权限、操作级鉴权、限流和审计日志。",
        "系统运维：提供Docker Compose与Caddy部署配置、健康检查、结构化日志和基础运行指标。",
        "扩展功能：在核心链路稳定后实现多窗格工作区、受限文件浏览、限时只读或读写分享。",
    ]
    for i, item in enumerate(requirements, 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "二 主要技术指标与提交成果", 1)
    metrics = [
        "Hub和Agent采用Go实现，Web端采用React、TypeScript、Vite、TanStack Query和xterm.js，数据持久化采用PostgreSQL。",
        "Agent只建立到Hub TCP 443端口的出站连接，不监听公网端口；Agent以普通系统用户权限运行。",
        "系统应完成Linux与macOS接入验证，至少验证3台异构主机、每台5个tmux会话和20路并发终端连接。",
        "稳定网络下按键到回显P95附加时延目标不高于200毫秒；短时断网恢复时间目标不高于10秒。",
        "设备删除或证书吊销后，既有连接和未使用票据应立即失效；越权测试请求拦截率达到100%。",
        "完成单元测试、集成测试、端到端测试、故障注入和安全测试，记录测试环境、步骤与结果。",
        "提交SinthMux软件成品、完整源代码、数据库迁移、部署脚本、接口文档、测试报告、用户手册、毕业论文和答辩材料。",
    ]
    for i, item in enumerate(metrics, 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "三 进度安排", 1)
    schedule = [
        ("2027.2.16—2027.3.1", "完成文献调研、需求分析、开题报告和总体架构设计，确定核心协议与数据模型。"),
        ("2027.3.2—2027.3.15", "完成工程初始化、Hub与Agent最小长连接、设备注册表和心跳机制。"),
        ("2027.3.16—2027.3.31", "完成远程tmux会话管理、终端流中继和浏览器xterm.js接入，准备中期检查。"),
        ("2027.4.1—2027.4.20", "完成一次性配对码、设备证书、短期票据、权限校验、审计日志和断线恢复。"),
        ("2027.4.21—2027.5.10", "完善响应式Web/PWA、多窗格工作区及受限文件功能，完成公网部署。"),
        ("2027.5.11—2027.5.24", "完成单元、集成、端到端、安全、性能和故障注入测试，整理实验数据。"),
        ("2027.5.25—2027.6.7", "完成论文初稿、系统使用手册、部署文档与测试报告，进行论文查重和修改。"),
        ("2027.6.8—2027.6.21", "准备演示环境、答辩PPT和测试数据，完成预答辩与正式答辩。"),
        ("2027.6.22—2027.6.28", "根据答辩意见修改论文和系统，归档全部毕业设计材料。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_text(table.rows[0].cells[0], "时间", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
    set_cell_text(table.rows[0].cells[1], "主要任务", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
    set_repeat_table_header(table.rows[0])
    for date, task in schedule:
        cells = table.add_row().cells
        set_cell_text(cells[0], date, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], task, size=9.5)
        prevent_row_split(table.rows[-1])

    doc.add_paragraph()
    add_signature_box(doc, "审批意见", "课题目标明确，技术路线合理，工作量符合软件工程专业本科毕业设计要求，同意下达任务。", "教研室主任：________________    2026年12月21日")

    path = OUT / f"1-{STUDENT_ID}-{STUDENT}-任务书-SinthMux.docx"
    doc.save(path)
    return path


def build_opening_report():
    doc = configure_doc(Document())
    # Cover
    doc.add_paragraph("\n\n")
    add_center_title(doc, [(SCHOOL, 22, True), ("本科毕业设计论文开题报告", 24, True)])
    doc.add_paragraph("\n")
    cover_rows = [
        ("学院", COLLEGE),
        ("专业班级", CLASS_NAME),
        ("设计题目", TITLE),
        ("学生姓名", STUDENT),
        ("学号", STUDENT_ID),
        ("起迄日期", "2027年3月1日—2027年6月28日"),
        ("设计地点", SCHOOL),
        ("指导教师", ADVISER),
    ]
    table = add_label_value_table(doc, cover_rows, widths=(4.0, 11.8))
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
    doc.add_paragraph("\n")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("2027年3月1日")
    set_run_font(r, 12)
    doc.add_page_break()

    add_center_title(doc, [("毕业设计论文开题报告", 18, True)])
    add_body(doc, "本报告围绕SinthMux跨设备共享持久终端系统展开。课题拟解决个人多台电脑和云服务器位于不同网络、缺少公网入口、终端会话难以跨设备连续使用等问题，采用公网Hub与远程Agent主动反向连接的方式，实现多主机集中管理、tmux持久会话、浏览器实时终端、断线恢复、细粒度权限和安全审计。", indent=True)

    add_heading(doc, "一 问题的提出和课题概述", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "远程开发和云计算使个人用户经常同时使用办公电脑、家庭电脑、实验室主机和云服务器。传统SSH要求控制端能够直接访问目标主机；当主机位于家庭路由器、校园网或企业防火墙之后时，用户通常需要公网IP、端口映射、VPN或跳板机，配置成本较高。即使连接建立，网络切换和终端关闭也可能造成前台任务中断，手机和平板上的连续操作体验尤其不足。")
    add_body(doc, "tmux能够将Shell进程与客户端连接解耦，使会话在用户退出后继续运行，但原生tmux主要面向单机命令行使用，缺少跨主机聚合、浏览器访问、设备身份、权限分享和审计能力。因此，本课题拟在tmux持久会话之上构建SinthMux，将分散主机统一接入公网Hub，并通过Web/PWA提供跨设备接续能力。")

    add_heading(doc, "1.2 研究目标与意义", 2)
    add_body(doc, "课题目标是设计并实现一个可公开部署的跨设备共享持久终端系统。远程Agent仅需主动访问Hub的TCP 443端口，不监听公网服务；Hub负责用户认证、设备注册、在线路由、终端中继、授权和审计；浏览器负责主机与会话管理及终端交互。该架构可以降低NAT和防火墙环境下的接入门槛，并将终端现场从单一设备扩展为可持续、可恢复和可受控共享的工作空间。")
    add_body(doc, "在软件工程方面，本课题覆盖需求分析、分布式架构、实时通信、身份认证、安全授权、数据库设计、响应式前端、自动化测试、容器部署和性能评价，具有完整系统设计与实现价值。在实验方面，可通过接入成功率、端到端交互时延、断线恢复时间、并发资源占用和越权拦截率对方案进行量化评价。")

    add_heading(doc, "1.3 主要研究内容", 2)
    contents = [
        "研究Agent主动反向连接Hub的多主机接入模型，设计设备注册、心跳、在线状态和路由机制。",
        "研究浏览器、Hub、Agent与tmux之间的终端流协议，实现输入、输出、尺寸调整、心跳、关闭和错误消息。",
        "研究一次性配对码、设备证书、短期终端票据、角色权限和操作审计等公网安全机制。",
        "实现响应式Web/PWA、多主机会话列表、实时终端和移动端特殊键，增强阶段实现多窗格与受限文件管理。",
        "建立自动化测试与实验环境，评价跨网络接入、终端时延、断线恢复、并发能力和安全控制效果。",
    ]
    for i, item in enumerate(contents, 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "二 国内外研究与相关技术综述", 1)
    add_heading(doc, "2.1 Web终端与持久会话", 2)
    add_body(doc, "Web终端通常通过浏览器终端模拟器、WebSocket和服务器端PTY实现。xterm.js在浏览器中负责ANSI终端显示和输入，ttyd与WeTTY等项目展示了将PTY或SSH会话映射到WebSocket的基本方法。此类项目部署简单、交互直接，但多数以单主机或单命令暴露为目标，缺少完整的多主机设备生命周期和集中授权模型。")
    add_body(doc, "tmux通过服务端维护会话、窗口和窗格，使任务不依赖某个SSH连接而长期运行。SinthMux不替代tmux，而是在其上增加设备注册、Hub路由、跨设备Web访问和安全控制，从而把单机持久会话扩展为多主机工作台。")

    add_heading(doc, "2.2 集中式远程访问平台", 2)
    add_body(doc, "ShellHub采用中央网关管理位于边缘与云环境中的设备，强调NAT后主机无需暴露公网SSH端口；MeshCentral通过Agent实现远程终端、文件和设备管理；Teleport提供基于短期证书、角色权限和审计的基础设施访问。这些系统验证了Agent主动接入控制面的可行性，但其目标分别偏向SSH网关、综合设备管理或企业零信任平台，功能体量明显超过本科毕业设计。")
    add_body(doc, "SinthMux吸收这些项目的连接模型和安全思想，将范围限定为个人及小团队多主机tmux工作区。这样既保留NAT适应性、统一身份和审计能力，又避免同时实现RDP、Kubernetes、数据库代理等无关协议。")

    add_heading(doc, "2.3 协作终端与连接连续性", 2)
    add_body(doc, "sshx与TermPair关注浏览器协作终端、自动重连、低延迟及可选端到端加密，对SinthMux的终端分享和恢复机制具有参考价值。WebSocket适合经TCP 443穿越常见代理和防火墙；TLS 1.3与mTLS可分别保障传输机密性和设备身份。对于带副作用的终端输入，应采用一次性短期票据和服务端操作授权，避免长期令牌出现在URL或浏览器本地存储中。")

    add_heading(doc, "2.4 现有方案不足与本课题切入点", 2)
    add_body(doc, "现有轻量Web终端缺少多主机Hub和设备身份，企业远程访问平台又存在功能庞大、部署复杂的问题。SinthMux选择“公网Hub、Agent主动出站、tmux持久会话、浏览器跨设备接续”作为核心组合，并把断线恢复、短期票据和操作审计纳入统一设计。创新点主要体现在面向个人多设备场景的工程组合、可解释安全边界和可量化的连续性评价，而非宣称提出新的密码算法或基础网络协议。")

    add_heading(doc, "三 参考文献", 1)
    for ref in REFERENCES:
        add_body(doc, ref, indent=False, size=9.5, after=2, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_page_break()
    add_heading(doc, "四 研究或解决的问题和拟采用的方法", 1)
    add_heading(doc, "4.1 拟解决的主要问题", 2)
    problems = [
        "多网络环境接入问题：远程电脑可能位于NAT或防火墙后，Hub无法主动建立连接。",
        "多主机统一路由问题：需要可靠维护设备身份、在线状态、能力、连接和目标tmux会话之间的映射。",
        "终端实时性与连续性问题：需要处理双向字节流、窗口尺寸、背压、心跳、半开连接和网络切换后的恢复。",
        "公网安全问题：需要防止伪造设备、票据重放、跨用户越权、WebSocket滥用和敏感日志泄漏。",
        "跨设备交互问题：桌面、手机和平板的输入方式和屏幕尺寸不同，需要针对移动端设计特殊键与可恢复界面。",
        "工程可验证问题：需要建立可重复的多主机测试环境和量化指标，验证方案的可靠性、性能和安全性。",
    ]
    for i, item in enumerate(problems, 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "4.2 系统总体架构", 2)
    add_body(doc, "系统采用浏览器、Hub、Agent、tmux四层结构。浏览器通过HTTPS和WSS访问Hub；Agent主动建立到Hub的mTLS WebSocket长连接；Hub维护在线设备注册表，并把浏览器操作路由至目标Agent；Agent使用结构化参数调用tmux并通过PTY转发终端字节流。远程主机不开放新增入站端口。")
    architecture = doc.add_table(rows=1, cols=3)
    architecture.style = "Table Grid"
    architecture.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["层次", "主要职责", "拟采用技术"]
    for i, h in enumerate(headers):
        set_cell_text(architecture.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
    set_repeat_table_header(architecture.rows[0])
    for values in [
        ("Web/PWA", "主机与会话管理、终端交互、移动端适配", "React、TypeScript、TanStack Query、xterm.js"),
        ("Hub", "认证、设备注册、路由、票据、审计、REST API", "Go、chi、coder/websocket、PostgreSQL"),
        ("Agent", "主动连接、心跳、tmux/PTY适配、受限文件访问", "Go、mTLS WebSocket、systemd/launchd"),
        ("基础设施", "TLS入口、容器部署、日志与指标", "Caddy、Docker Compose、Prometheus指标"),
    ]:
        cells = architecture.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
        prevent_row_split(architecture.rows[-1])

    add_heading(doc, "4.3 关键方法", 2)
    methods = [
        "设备配对与身份：Hub生成5分钟有效的一次性配对码；Agent本机生成P-256密钥和CSR；Hub验证后签发短期设备证书，并通过证书SAN绑定用户空间与deviceId。",
        "连接与心跳：Agent通过mTLS WebSocket主动连接Hub，每15秒发送心跳；Hub连续45秒未收到心跳则标记离线；Agent按指数退避重连。",
        "终端票据与中继：浏览器申请30至60秒有效且只能兑换一次的终端票据，Hub校验用户、设备、会话和权限后建立WSS终端流。",
        "协议复用：Agent连接使用版本化Envelope，区分请求、响应、流打开、流数据、流关闭和事件，并通过streamId复用多个终端。",
        "权限与审计：定义Owner、Admin、Operator、Viewer角色，对会话查看、输入、创建、关闭和设备管理执行服务端鉴权，记录关键操作但不记录完整终端输入。",
        "路径与命令安全：tmux调用全部使用结构化参数，不拼接Shell字符串；文件访问限制在配置根目录，并检查规范化路径与符号链接。",
    ]
    for i, item in enumerate(methods, 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "4.4 数据设计", 2)
    add_body(doc, "核心实体包括users、spaces、memberships、devices、device_certificates、pairing_codes、sessions、terminal_tickets、audit_events和user_settings。设备以服务端UUID标识，不以IP或名称作为身份；会话以device_id、runtime和runtime_ref联合定位；票据只保存哈希、目标、有效期和使用状态；审计事件采用追加写。")

    add_heading(doc, "4.5 测试与实验方法", 2)
    experiments = [
        "功能测试：覆盖设备配对、上线离线、tmux会话增删改查、终端输入输出、窗口尺寸和权限控制。",
        "可靠性测试：模拟Agent进程退出、Hub重启、网络中断、心跳超时和移动网络切换，记录恢复时间与消息完整性。",
        "性能测试：测量终端首屏时间、按键到回显P50/P95时延、20至100路并发连接下Hub的CPU与内存。",
        "安全测试：验证无效证书、过期票据、重复票据、越权设备访问、非法Origin、超大消息和路径穿越均被拒绝。",
        "对照实验：比较逐台SSH、无恢复Web终端和SinthMux在接入步骤、会话恢复时间和任务持续性方面的差异。",
    ]
    for i, item in enumerate(experiments, 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "五 可行性分析与预期成果", 1)
    add_heading(doc, "5.1 技术可行性", 2)
    add_body(doc, "Go具备成熟的HTTP、TLS、WebSocket和并发支持，适合构建Hub与跨平台Agent；React和xterm.js可实现浏览器终端；tmux提供稳定的持久会话基础；PostgreSQL、Caddy和Docker Compose能够满足数据存储与公网部署需求。ShellHub、MeshCentral、ttyd和sshx等开源项目已分别验证Agent反向连接、设备管理和Web终端的可行性。")
    add_heading(doc, "5.2 工作量可行性", 2)
    add_body(doc, "课题将范围限定为tmux持久终端，不实现RDP、VNC、Kubernetes和数据库代理。开发按“远程链路验证、tmux产品闭环、安全与部署、增强体验、实验与论文”五阶段推进，核心功能和扩展功能分层验收，可在本科毕业设计周期内完成。")
    add_heading(doc, "5.3 预期成果", 2)
    for i, item in enumerate([
        "可公网部署的SinthMux Hub、Linux/macOS Agent和响应式Web/PWA。",
        "支持至少3台异构主机和20路并发终端的可演示系统。",
        "完整源代码、数据库迁移、部署脚本、接口文档、测试报告和用户手册。",
        "接入成功率、交互时延、断线恢复、资源占用和安全测试实验数据。",
        "符合学校规范的毕业设计论文和答辩演示材料。",
    ], 1):
        add_numbered(doc, f"{i}）", item)

    add_heading(doc, "六 进度安排", 1)
    schedule = [
        ("2027.3.1—3.15", "完成需求分析、协议设计、Hub与Agent最小长连接。"),
        ("2027.3.16—3.31", "完成设备注册、心跳、tmux会话管理和终端中继。"),
        ("2027.4.1—4.20", "完成配对码、mTLS、短期票据、权限与审计。"),
        ("2027.4.21—5.10", "完成响应式Web/PWA、多窗格、受限文件功能和公网部署。"),
        ("2027.5.11—5.24", "完成自动化测试、故障注入、安全测试和性能实验。"),
        ("2027.5.25—6.7", "完成论文初稿、用户手册、部署文档和查重修改。"),
        ("2027.6.8—6.21", "完成演示数据、答辩PPT、预答辩和正式答辩。"),
        ("2027.6.22—6.28", "根据答辩意见修改并归档系统和论文。"),
    ]
    schedule_table = doc.add_table(rows=1, cols=2)
    schedule_table.style = "Table Grid"
    set_cell_text(schedule_table.rows[0].cells[0], "时间", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
    set_cell_text(schedule_table.rows[0].cells[1], "工作内容", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font=HEAD_FONT)
    set_repeat_table_header(schedule_table.rows[0])
    for date, task in schedule:
        cells = schedule_table.add_row().cells
        set_cell_text(cells[0], date, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], task, size=9.5)
        prevent_row_split(schedule_table.rows[-1])

    doc.add_paragraph()
    add_signature_box(doc, "指导教师意见", "该课题面向跨网络、跨设备的远程终端持续使用问题，研究目标明确，技术路线完整，涵盖分布式连接、实时通信、身份认证、权限控制、响应式前端和系统测试等内容，具有一定工程深度与综合工作量。课题范围控制合理，预期成果和实验指标明确，具备实施条件。综上，同意开题。", "指导教师：________________    2027年3月3日")
    doc.add_paragraph()
    add_signature_box(doc, "审批意见", "课题内容符合软件工程专业培养目标，工作计划合理，同意开题。", "教研室主任：________________    2027年3月4日")

    path = OUT / f"2-{STUDENT_ID}-{STUDENT}-开题报告-SinthMux.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for output in (build_application(), build_task_book(), build_opening_report()):
        print(output)
